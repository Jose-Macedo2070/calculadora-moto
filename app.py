
import json
import os
from datetime import date, datetime, time, timedelta
from io import BytesIO

import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# =========================================================
# 1. CONFIGURAÇÕES
# =========================================================

st.set_page_config(
    page_title="Finanças Bajaj",
    page_icon="🏍️",
    layout="centered",
    initial_sidebar_state="collapsed",
)

ARQUIVO_BD = "banco_entregas_bajaj.json"

META_BRUTA_SEMANAL = 1000.00
META_CUSTOS_SEMANAL = 405.00

TAXA_RETENCAO = 0.405
TAXA_LUCRO = 0.595

PERCENTUAL_PAI = 0.3079
PERCENTUAL_GASOLINA = 0.2222
PERCENTUAL_DESPESAS = 0.4699

COLUNAS_VALORES = [
    "Bruto (R$)",
    "Lucro Livre (R$)",
    "Pai (R$)",
    "Gasolina (R$)",
    "Despesas Gerais (R$)",
]

COLUNAS_HISTORICO = ["Data", "Hora"] + COLUNAS_VALORES


# =========================================================
# 2. ESTILO RESPONSIVO PARA CELULAR
# =========================================================

st.markdown(
    """
    <style>
    .block-container {
        max-width: 1000px;
        padding-top: 1.2rem;
        padding-bottom: 2rem;
        padding-left: 1rem;
        padding-right: 1rem;
    }

    div[data-testid="stMetric"] {
        background: rgba(128, 128, 128, 0.07);
        border: 1px solid rgba(128, 128, 128, 0.15);
        padding: 12px;
        border-radius: 12px;
        overflow-wrap: anywhere;
    }

    div[data-testid="stMetricValue"] {
        font-size: clamp(1.15rem, 4vw, 1.8rem);
    }

    div.stButton > button,
    div.stDownloadButton > button,
    div.stFormSubmitButton > button {
        width: 100%;
        min-height: 44px;
        border-radius: 10px;
        white-space: normal;
    }

    div[data-testid="stProgress"] {
        margin: 8px 0 14px 0;
    }

    /* Evita que várias colunas fiquem apertadas no celular. */
    @media (max-width: 640px) {
        .block-container {
            padding-left: 0.75rem;
            padding-right: 0.75rem;
            padding-top: 0.8rem;
        }

        div[data-testid="stHorizontalBlock"] {
            flex-wrap: wrap;
            gap: 0.65rem;
        }

        div[data-testid="stHorizontalBlock"] > div {
            min-width: 100% !important;
            flex: 1 1 100% !important;
        }

        h1 {
            font-size: 1.65rem !important;
        }

        h2 {
            font-size: 1.35rem !important;
        }

        h3 {
            font-size: 1.15rem !important;
        }

        div[data-testid="stMetric"] {
            padding: 10px;
        }

        /* A tabela pode ser deslizada horizontalmente. */
        div[data-testid="stDataFrame"] {
            max-width: 100%;
        }
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# 3. FUNÇÕES GERAIS
# =========================================================

def moeda(valor):
    """Formata valores no padrão brasileiro."""

    texto = f"{float(valor):,.2f}"
    texto = (
        texto.replace(",", "X")
        .replace(".", ",")
        .replace("X", ".")
    )
    return f"R$ {texto}"


def carregar_banco():
    """Lê o histórico sem apagar arquivos com problemas."""

    if not os.path.exists(ARQUIVO_BD):
        return []

    try:
        with open(ARQUIVO_BD, "r", encoding="utf-8") as arquivo:
            dados = json.load(arquivo)

        if not isinstance(dados, list):
            raise ValueError("O arquivo não contém uma lista de registros.")

        return dados

    except (json.JSONDecodeError, OSError, ValueError) as erro:
        st.error(
            f"Não foi possível carregar o histórico: {erro}. "
            "Verifique o arquivo JSON antes de continuar."
        )
        st.stop()


def salvar_banco(dados):
    """Grava os dados usando um arquivo temporário."""

    temporario = ARQUIVO_BD + ".tmp"

    try:
        with open(temporario, "w", encoding="utf-8") as arquivo:
            json.dump(
                dados,
                arquivo,
                indent=4,
                ensure_ascii=False,
            )

        os.replace(temporario, ARQUIVO_BD)

    except OSError as erro:
        if os.path.exists(temporario):
            try:
                os.remove(temporario)
            except OSError:
                pass

        st.error(f"Não foi possível salvar os dados: {erro}")
        st.stop()


# =========================================================
# 4. CÁLCULOS FINANCEIROS
# =========================================================

def calcular_registro(valor_bruto, data_corrida, hora_corrida):
    """Calcula a distribuição financeira da corrida."""

    valor_bruto = round(float(valor_bruto), 2)
    retencao = valor_bruto * TAXA_RETENCAO

    if isinstance(hora_corrida, time):
        hora_texto = hora_corrida.strftime("%H:%M")
    else:
        hora_texto = str(hora_corrida or "")

    return {
        "Data": data_corrida.strftime("%d/%m/%Y"),
        "Hora": hora_texto,
        "Bruto (R$)": valor_bruto,
        "Lucro Livre (R$)": round(
            valor_bruto * TAXA_LUCRO, 2
        ),
        "Pai (R$)": round(
            retencao * PERCENTUAL_PAI, 2
        ),
        "Gasolina (R$)": round(
            retencao * PERCENTUAL_GASOLINA, 2
        ),
        "Despesas Gerais (R$)": round(
            retencao * PERCENTUAL_DESPESAS, 2
        ),
    }


def normalizar_historico(dados):
    """
    Mantém a compatibilidade com registros antigos.
    Horários inexistentes permanecem vazios.
    """

    resultado = []

    for registro in dados:
        try:
            data_corrida = datetime.strptime(
                registro.get("Data", ""),
                "%d/%m/%Y",
            ).date()
        except (ValueError, TypeError):
            raise ValueError(
                "Foi encontrado um registro com data inválida."
            )

        try:
            bruto = float(registro.get("Bruto (R$)", 0) or 0)
        except (ValueError, TypeError):
            raise ValueError(
                "Foi encontrado um registro com valor inválido."
            )

        hora_texto = str(registro.get("Hora", "") or "")

        try:
            hora_texto = datetime.strptime(
                hora_texto, "%H:%M"
            ).strftime("%H:%M") if hora_texto else ""
        except ValueError:
            hora_texto = ""

        novo = calcular_registro(
            bruto,
            data_corrida,
            hora_texto,
        )

        # Preserva os valores financeiros já registrados.
        for coluna in COLUNAS_VALORES:
            if coluna in registro:
                try:
                    novo[coluna] = float(registro[coluna] or 0)
                except (ValueError, TypeError):
                    raise ValueError(
                        f"Valor inválido na coluna {coluna}."
                    )

        # Compatibilidade com categorias antigas.
        if (
            "Despesas Gerais (R$)" not in registro
            and (
                "Manutenção (R$)" in registro
                or "Seguro/Reserva (R$)" in registro
            )
        ):
            novo["Despesas Gerais (R$)"] = round(
                float(registro.get("Manutenção (R$)", 0) or 0)
                + float(registro.get("Seguro/Reserva (R$)", 0) or 0),
                2,
            )

        resultado.append(novo)

    return resultado


def obter_semana_atual():
    hoje = date.today()
    inicio = hoje - timedelta(days=hoje.weekday())
    fim = inicio + timedelta(days=6)
    return inicio, fim


# =========================================================
# 5. EXPORTAÇÃO EXCEL
# =========================================================

def gerar_excel(dataframe):
    buffer = BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        dataframe.to_excel(
            writer,
            index=False,
            sheet_name="Entregas",
        )

        planilha = writer.sheets["Entregas"]

        fundo = PatternFill(
            start_color="1B365D",
            end_color="1B365D",
            fill_type="solid",
        )

        fonte_cabecalho = Font(
            name="Arial",
            size=11,
            bold=True,
            color="FFFFFF",
        )

        fonte_dados = Font(
            name="Arial",
            size=10,
            color="000000",
        )

        lado = Side(style="thin", color="D3D3D3")

        borda = Border(
            left=lado,
            right=lado,
            top=lado,
            bottom=lado,
        )

        for celula in planilha[1]:
            celula.fill = fundo
            celula.font = fonte_cabecalho
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            celula.border = borda

        for linha in planilha.iter_rows(min_row=2):
            for celula in linha:
                celula.font = fonte_dados
                celula.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )
                celula.border = borda

                nome_coluna = planilha.cell(
                    row=1,
                    column=celula.column,
                ).value

                if nome_coluna in COLUNAS_VALORES:
                    celula.number_format = '"R$" #,##0.00'

        for coluna in planilha.columns:
            tamanho = max(
                len(str(celula.value or ""))
                for celula in coluna
            )

            letra = get_column_letter(coluna[0].column)
            planilha.column_dimensions[letra].width = max(
                tamanho + 4, 12
            )

        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions
        planilha.row_dimensions[1].height = 25

    buffer.seek(0)
    return buffer


# =========================================================
# 6. CARREGAR HISTÓRICO
# =========================================================

try:
    historico = normalizar_historico(carregar_banco())
except (ValueError, TypeError, KeyError) as erro:
    st.error(
        f"O histórico possui registros incompatíveis: {erro}. "
        "Faça uma cópia do JSON antes de corrigir os dados."
    )
    st.stop()


# =========================================================
# 7. CABEÇALHO
# =========================================================

st.title("🏍️ Gestão Financeira Bajaj")
st.caption(
    "Dominar 200 | Meta semanal: R$ 1.000,00 "
    "| Retenção planejada: 40,5%"
)


# =========================================================
# 8. NAVEGAÇÃO
# =========================================================

TELAS = [
    "Nova corrida",
    "Dashboard",
    "Histórico",
    "Gerenciar",
]

if "tela_atual" not in st.session_state:
    st.session_state.tela_atual = "Dashboard"

st.markdown("### Menu")

# O layout de duas colunas no desktop vira uma coluna
# por botão em telas pequenas, por causa do CSS responsivo.
coluna_esquerda, coluna_direita = st.columns(2)

with coluna_esquerda:
    if st.button(
        "➕ Nova corrida",
        type=(
            "primary"
            if st.session_state.tela_atual == "Nova corrida"
            else "secondary"
        ),
        use_container_width=True,
    ):
        st.session_state.tela_atual = "Nova corrida"
        st.rerun()

    if st.button(
        "📊 Dashboard",
        type=(
            "primary"
            if st.session_state.tela_atual == "Dashboard"
            else "secondary"
        ),
        use_container_width=True,
    ):
        st.session_state.tela_atual = "Dashboard"
        st.rerun()

with coluna_direita:
    if st.button(
        "📋 Histórico",
        type=(
            "primary"
            if st.session_state.tela_atual == "Histórico"
            else "secondary"
        ),
        use_container_width=True,
    ):
        st.session_state.tela_atual = "Histórico"
        st.rerun()

    if st.button(
        "⚙️ Gerenciar",
        type=(
            "primary"
            if st.session_state.tela_atual == "Gerenciar"
            else "secondary"
        ),
        use_container_width=True,
    ):
        st.session_state.tela_atual = "Gerenciar"
        st.rerun()

st.divider()


# =========================================================
# 9. TELA NOVA CORRIDA
# =========================================================

if st.session_state.tela_atual == "Nova corrida":

    st.subheader("📥 Lançar Nova Corrida")

    st.write(
        "Informe o valor recebido, a data e o horário."
    )

    with st.form("form_nova_corrida", clear_on_submit=True):

        valor_bruto = st.number_input(
            "Valor bruto recebido (R$)",
            min_value=0.0,
            step=5.0,
            format="%.2f",
        )

        data_corrida = st.date_input(
            "Data da corrida",
            value=date.today(),
            format="DD/MM/YYYY",
        )

        hora_corrida = st.time_input(
            "Horário da corrida",
            value=datetime.now().time().replace(
                second=0,
                microsecond=0,
            ),
            step=timedelta(minutes=5),
            format="HH:mm",
        )

        salvar = st.form_submit_button(
            "💾 Calcular e Salvar Corrida",
            use_container_width=True,
        )

        if salvar:
            if valor_bruto <= 0:
                st.warning("Informe um valor maior que zero.")
            else:
                novo = calcular_registro(
                    valor_bruto,
                    data_corrida,
                    hora_corrida,
                )

                historico.append(novo)
                salvar_banco(historico)

                st.success("Corrida registrada com sucesso!")
                st.rerun()

    st.info(
        "O sistema calculará automaticamente seu lucro livre "
        "e a divisão entre Pai, Gasolina e Despesas Gerais."
    )


# =========================================================
# 10. TELA DASHBOARD
# =========================================================

elif st.session_state.tela_atual == "Dashboard":

    st.subheader("📊 Painel Financeiro")

    if not historico:
        st.info(
            "Nenhuma corrida registrada. "
            "Acesse 'Nova corrida' para começar."
        )

    else:
        df = pd.DataFrame(
            historico,
            columns=COLUNAS_HISTORICO,
        )

        bruto_total = df["Bruto (R$)"].sum()
        lucro_total = df["Lucro Livre (R$)"].sum()

        pai_total = df["Pai (R$)"].sum()
        gasolina_total = df["Gasolina (R$)"].sum()
        despesas_total = df["Despesas Gerais (R$)"].sum()

        df["_Data"] = pd.to_datetime(
            df["Data"],
            format="%d/%m/%Y",
            errors="coerce",
        )

        inicio, fim = obter_semana_atual()

        semana = df[
            (df["_Data"].dt.date >= inicio)
            & (df["_Data"].dt.date <= fim)
        ]

        bruto_semana = semana["Bruto (R$)"].sum()
        lucro_semana = semana["Lucro Livre (R$)"].sum()

        custos_semana = (
            semana["Pai (R$)"].sum()
            + semana["Gasolina (R$)"].sum()
            + semana["Despesas Gerais (R$)"].sum()
        )

        st.caption(
            f"Semana atual: {inicio.strftime('%d/%m/%Y')} "
            f"a {fim.strftime('%d/%m/%Y')}"
        )

        c1, c2 = st.columns(2)

        c1.metric(
            "Faturamento total",
            moeda(bruto_total),
        )

        c2.metric(
            "Lucro livre total",
            moeda(lucro_total),
        )

        st.markdown("### 🔒 Caixinhas Reservadas")

        c1, c2, c3 = st.columns(3)

        c1.metric("👴 Pai", moeda(pai_total))
        c2.metric("⛽ Gasolina", moeda(gasolina_total))
        c3.metric("🛠️ Despesas", moeda(despesas_total))

        st.divider()

        st.markdown("### 🚀 Meta de Faturamento Semanal")

        progresso_bruto = min(
            bruto_semana / META_BRUTA_SEMANAL,
            1.0,
        )

        st.progress(progresso_bruto)

        st.metric(
            "Faturamento da semana",
            moeda(bruto_semana),
        )

        st.metric(
            "Falta para atingir a meta",
            moeda(max(META_BRUTA_SEMANAL - bruto_semana, 0)),
        )

        st.caption(
            f"Progresso: {progresso_bruto:.1%} "
            f"| Meta: {moeda(META_BRUTA_SEMANAL)}"
        )

        st.markdown("### 📈 Meta de Retenção e Custos")

        progresso_custos = min(
            custos_semana / META_CUSTOS_SEMANAL,
            1.0,
        )

        st.progress(progresso_custos)

        st.metric(
            "Reservado nesta semana",
            moeda(custos_semana),
        )

        st.metric(
            "Falta reservar",
            moeda(max(META_CUSTOS_SEMANAL - custos_semana, 0)),
        )

        st.caption(
            f"Progresso: {progresso_custos:.1%} "
            f"| Meta: {moeda(META_CUSTOS_SEMANAL)}"
        )

        st.metric(
            "Lucro livre da semana",
            moeda(lucro_semana),
        )

        if custos_semana >= META_CUSTOS_SEMANAL:
            st.success("Meta semanal de custos atingida!")


# =========================================================
# 11. TELA HISTÓRICO E EXPORTAÇÃO
# =========================================================

elif st.session_state.tela_atual == "Histórico":

    st.subheader("📋 Histórico de Corridas")

    if not historico:
        st.info("Nenhum lançamento encontrado.")

    else:
        df = pd.DataFrame(
            historico,
            columns=COLUNAS_HISTORICO,
        )

        df_ordenado = df.copy()

        df_ordenado["_DataHora"] = pd.to_datetime(
            df_ordenado["Data"]
            + " "
            + df_ordenado["Hora"].replace("", "00:00"),
            format="%d/%m/%Y %H:%M",
            errors="coerce",
        )

        df_ordenado = df_ordenado.sort_values(
            "_DataHora",
            ascending=False,
            na_position="last",
        )

        df_exibicao = df_ordenado[COLUNAS_HISTORICO]

        st.dataframe(
            df_exibicao,
            use_container_width=True,
            hide_index=True,
            height=400,
            column_config={
                "Data": st.column_config.TextColumn("Data"),
                "Hora": st.column_config.TextColumn("Hora"),
                **{
                    coluna: st.column_config.NumberColumn(
                        coluna,
                        format="R$ %.2f",
                    )
                    for coluna in COLUNAS_VALORES
                },
            },
        )

        st.caption(
            "No celular, deslize a tabela horizontalmente "
            "para consultar todas as colunas."
        )

        # O botão fica diretamente abaixo da tabela.
        st.markdown("#### 📥 Exportar Histórico")

        excel = gerar_excel(df_exibicao)

        st.download_button(
            label="📊 Baixar Planilha Excel",
            data=excel,
            file_name=(
                f"financeiro_bajaj_"
                f"{datetime.now().strftime('%m_%Y')}.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True,
        )


# =========================================================
# 12. TELA GERENCIAR
# =========================================================

elif st.session_state.tela_atual == "Gerenciar":

    st.subheader("⚙️ Gerenciar Lançamentos")

    if not historico:
        st.info("Não existem lançamentos para gerenciar.")

    else:
        opcoes = [
            (
                f"ID {i} | {registro['Data']} "
                f"{registro.get('Hora', '')} | "
                f"{moeda(registro['Bruto (R$)'])}"
            )
            for i, registro in enumerate(historico)
        ]

        selecionado = st.selectbox(
            "Selecione uma corrida",
            options=opcoes,
        )

        indice = int(
            selecionado.split("|")[0].replace("ID", "").strip()
        )

        registro = historico[indice]

        st.markdown("### ✏️ Editar Corrida")

        data_original = datetime.strptime(
            registro["Data"],
            "%d/%m/%Y",
        ).date()

        try:
            hora_original = datetime.strptime(
                registro.get("Hora", ""),
                "%H:%M",
            ).time()
        except ValueError:
            hora_original = time(0, 0)

        with st.form("form_editar_corrida"):

            valor_editado = st.number_input(
                "Valor bruto (R$)",
                min_value=0.0,
                value=float(registro["Bruto (R$)"]),
                step=5.0,
                format="%.2f",
            )

            data_editada = st.date_input(
                "Data",
                value=data_original,
                format="DD/MM/YYYY",
                key="data_edicao",
            )

            hora_editada = st.time_input(
                "Horário",
                value=hora_original,
                step=timedelta(minutes=5),
                format="HH:mm",
                key="hora_edicao",
            )

            confirmar = st.form_submit_button(
                "Salvar Alterações",
                use_container_width=True,
            )

            if confirmar:
                if valor_editado <= 0:
                    st.warning("O valor deve ser maior que zero.")
                else:
                    historico[indice] = calcular_registro(
                        valor_editado,
                        data_editada,
                        hora_editada,
                    )

                    salvar_banco(historico)

                    st.success("Corrida atualizada!")
                    st.rerun()

        st.divider()

        st.markdown("### 🗑️ Excluir Corrida")

        confirmar_exclusao = st.checkbox(
            "Confirmo que desejo excluir somente esta corrida."
        )

        if st.button(
            "Excluir Corrida Selecionada",
            disabled=not confirmar_exclusao,
            use_container_width=True,
        ):
            historico.pop(indice)
            salvar_banco(historico)

            st.success("Corrida excluída!")
            st.rerun()

        st.divider()

    st.markdown("### 🚨 Limpar Todo o Histórico")

    st.warning(
        "Esta ação remove permanentemente todos os registros "
        "armazenados no arquivo JSON."
    )

    confirmar_tudo = st.checkbox(
        "Confirmo que desejo apagar todo o histórico."
    )

    if st.button(
        "💥 Apagar Todo o Histórico",
        disabled=not confirmar_tudo or not historico,
        use_container_width=True,
    ):
        salvar_banco([])
        st.success("Histórico apagado!")
        st.rerun()
